#!/usr/bin/env python3
"""compare_xlsx_rendered.py - compare two workbooks (or two directories of
workbooks) cell by cell on what Excel renders: value, font name/size/bold/
italic/color, solid fill color, the four borders, horizontal/vertical
alignment and wrap, plus merged ranges, column widths and row heights.

Encodings that render identically compare equal: an absent border and
style="none", horizontal="general" and no alignment, vertical="bottom" and
none.  Usage: compare_xlsx_rendered.py OLD NEW [--result-file F]
Prints TOTAL_DIFFS n; with --result-file writes PASS/FAIL there.
"""
import sys, glob, os
from openpyxl import load_workbook
def norm_color(c):
    if c is None: return None
    v = c.rgb if isinstance(c.rgb, str) else None
    if v is None: return None
    v = v.upper()
    return v[-6:] if len(v) >= 6 else v
def cell_sig(c):
    f = c.font; fl = c.fill; b = c.border; a = c.alignment
    fill = norm_color(fl.fgColor) if fl is not None and fl.fill_type == "solid" else None
    bs = lambda s: (s.style if s is not None and s.style not in (None, "none") else None)
    return (c.value, f.name, float(f.sz) if f.sz else None, bool(f.b), bool(f.i), norm_color(f.color) if f.color is not None else None,
            fill, bs(b.left), bs(b.right), bs(b.top), bs(b.bottom),
            a.horizontal if a.horizontal not in (None, "general") else None,
            a.vertical if a.vertical not in (None, "bottom") else None, bool(a.wrap_text))
args = [a for a in sys.argv[1:]]
result_file = None
if "--result-file" in args:
    i = args.index("--result-file"); result_file = args[i + 1]; del args[i:i + 2]
old_dir, new_dir = args[0], args[1]
if os.path.isdir(old_dir):
    pairs = [(po, os.path.join(new_dir, os.path.basename(po))) for po in sorted(glob.glob(os.path.join(old_dir, "*.xlsx")))]
else:
    pairs = [(old_dir, new_dir)]
bad = 0
for po, pn in pairs:
    wo, wn = load_workbook(po), load_workbook(pn)
    if wo.sheetnames != wn.sheetnames:
        print("SHEETS", po, wo.sheetnames, wn.sheetnames); bad += 1; continue
    for sn in wo.sheetnames:
        so, snw = wo[sn], wn[sn]
        mr = max(so.max_row, snw.max_row); mc = max(so.max_column, snw.max_column)
        diffs = 0
        for r in range(1, mr + 1):
            for c in range(1, mc + 1):
                a, b = cell_sig(so.cell(r, c)), cell_sig(snw.cell(r, c))
                if a != b:
                    diffs += 1
                    if diffs <= 3: print("DIFF", os.path.basename(po), sn, so.cell(r, c).coordinate, "\n  old", a, "\n  new", b)
        mo = sorted(str(x) for x in so.merged_cells.ranges); mn = sorted(str(x) for x in snw.merged_cells.ranges)
        if mo != mn: print("MERGE", os.path.basename(po), sn, mo, mn); diffs += 1
        wo_ = {k: v.width for k, v in so.column_dimensions.items()}; wn_ = {k: v.width for k, v in snw.column_dimensions.items()}
        if wo_ != wn_: print("WIDTH", os.path.basename(po), sn, wo_, wn_); diffs += 1
        ho = {k: v.height for k, v in so.row_dimensions.items() if v.height}; hn = {k: v.height for k, v in snw.row_dimensions.items() if v.height}
        if ho != hn: print("HEIGHT", os.path.basename(po), sn); diffs += 1
        print(f"{os.path.basename(po)}[{sn}] {mr}x{mc}: {'OK' if diffs == 0 else str(diffs)+' diffs'}")
        bad += diffs
print("TOTAL_DIFFS", bad)
if result_file:
    with open(result_file, "w") as fh:
        fh.write("PASS\n" if bad == 0 else f"FAIL {bad}\n")
