"""Dump layout facts of one worksheet, one fact per line, for Stata to read.

Usage:
    python3 xlsx_facts.py BOOK.xlsx SHEET RESULT.txt

RESULT.txt receives, in this order:
    merge <range>                   every merged range, e.g. "merge B6:C6"
    width <col> <width>             every column with an explicit width
    bottom <cell> <style>           every cell with a bottom border
    value <cell> <text>             every non-empty cell (newlines as \\n)
A Stata suite then asserts presence or absence of exact lines. Reading the
workbook back with openpyxl keeps the oracle independent of the Mata xl()
writer that produced it.
"""

import sys

from openpyxl import load_workbook


def main(argv):
    if len(argv) != 4:
        print(__doc__)
        return 2
    book, sheet, result = argv[1:]
    ws = load_workbook(book)[sheet]
    lines = []
    for rng in sorted(str(r) for r in ws.merged_cells.ranges):
        lines.append(f"merge {rng}")
    for col, dim in sorted(ws.column_dimensions.items()):
        if dim.width is not None and dim.customWidth:
            lines.append(f"width {col} {round(float(dim.width), 2):g}")
    for row in ws.iter_rows():
        for cell in row:
            style = cell.border.bottom.style if cell.border is not None else None
            if style:
                lines.append(f"bottom {cell.coordinate} {style}")
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value) != "":
                text = str(cell.value).replace("\n", "\\n")
                lines.append(f"value {cell.coordinate} {text}")
    with open(result, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
