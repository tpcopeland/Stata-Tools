"""Write the used range of one worksheet, "rows cols", for Stata to read.

Usage:
    python3 xlsx_dims.py BOOK.xlsx SHEET RESULT.txt

The size is openpyxl's max_row x max_column: every row and column that holds
a cell in the sheet XML, title and note rows and the spacer column A
included. Reading it with openpyxl keeps the oracle independent of the
import excel, describe call that tabtools uses for its completion line.
"""
import sys

from openpyxl import load_workbook


def main(book, sheet, result):
    wb = load_workbook(book)
    names = {name.lower(): name for name in wb.sheetnames}
    ws = wb[names[sheet.lower()]]
    with open(result, "w", encoding="utf-8") as fh:
        fh.write(f"{ws.max_row} {ws.max_column}\n")


if __name__ == "__main__":
    main(*sys.argv[1:4])
