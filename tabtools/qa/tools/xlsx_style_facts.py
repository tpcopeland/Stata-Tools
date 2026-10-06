"""Dump cell style facts of one worksheet, one fact per line, for Stata to read.

Usage:
    python3 xlsx_style_facts.py BOOK.xlsx SHEET RESULT.txt

RESULT.txt receives, in this order:
    sheets <name> <name> ...        every sheet name of the workbook
    value <cell> <text>             every non-empty cell (newlines as \\n)
    font <cell> <name> <size>       font family and point size of every non-empty cell
    fill <cell> <rgb>               every cell with a solid fill (ARGB as stored)
    border <cell> <side> <style>    every border side that has a style
The workbook is read back with openpyxl, so the oracle is independent of the
Mata xl() writer that produced it. A Stata suite asserts the presence or the
absence of exact lines.
"""

import sys

from openpyxl import load_workbook


def main(argv):
    if len(argv) != 4:
        print(__doc__)
        return 2
    book, sheet, out = argv[1:]
    wb = load_workbook(book)
    lines = ["sheets " + " ".join(wb.sheetnames)]
    ws = wb[sheet]
    fonts, fills, borders = [], [], []
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value) != "":
                text = str(c.value).replace("\n", "\\n")
                lines.append(f"value {c.coordinate} {text}")
                fonts.append(f"font {c.coordinate} {c.font.name} {int(c.font.sz)}")
            if c.fill is not None and c.fill.fill_type == "solid":
                fills.append(f"fill {c.coordinate} {c.fill.fgColor.rgb}")
            for side in ("top", "bottom", "left", "right"):
                st = getattr(c.border, side).style
                if st:
                    borders.append(f"border {c.coordinate} {side} {st}")
    with open(out, "w") as fh:
        fh.write("\n".join(lines + fonts + fills + borders) + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
