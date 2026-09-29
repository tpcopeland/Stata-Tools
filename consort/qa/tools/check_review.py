# Author: Timothy P Copeland, Karolinska Institutet
# Regression artifact oracle, 2026-09-29
import csv, importlib.util, sys
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
root, mode, out, *extra = sys.argv[1:]
out = Path(out)
if mode == "fixture":
    with (out / "back.csv").open("w", newline="") as f:
        w = csv.writer(f); w.writerow(["label", "n", "remaining"])
        w.writerow(["Population", 10, ""])
        w.writerow(['Excluded,\n"reason"', 2, 'Adults,\n"eligible"'])
    sys.exit(0)
if mode == "snapshot":
    (out / "before.bin").write_bytes((out / "back.csv").read_bytes())
    sys.exit(0)
if mode == "unchanged":
    assert (out / "before.bin").read_bytes() == (out / "back.csv").read_bytes()
    print("ARTIFACT_CONTENT_PASS", mode)
    sys.exit(0)
spec = importlib.util.spec_from_file_location("diagram", Path(root) / "consort_diagram.py")
m = importlib.util.module_from_spec(spec); sys.modules["diagram"] = m; spec.loader.exec_module(m)
d = m.from_csv(str(out / "back.csv"))
with (out / "resolved.csv").open(newline="") as f:
    table = list(csv.DictReader(f))
with (out / "back.csv").open(newline="") as f:
    raw = list(csv.DictReader(f))
expected_n = int(extra[0]) if mode == "random" else 8
assert int(table[-1]["n_remaining"]) == expected_n
assert d.initial_n - sum(x[1] for x in d.exclusions) == expected_n
if mode == "quotes":
    assert raw[0]["label"] == 'All "eligible" patients'
    assert raw[1]["label"] == 'Not "eligible"'
    assert raw[-1]["remaining"] == 'Final "eligible" patients'
if mode == "multiline":
    assert raw[1]["label"] == 'Excluded,\n"reason"'
    assert raw[-1]["remaining"] == 'Analysis "quoted"'
if mode == "labels":
    expected = " ".join(extra)
    assert table[-1]["cohort_label"] == expected
    assert d.exclusions[-1][2] == expected
    if len(table) > 2:
        assert table[1]["cohort_label"] == "Milestone"
        assert d.exclusions[0][2] == "Milestone"
# Workbook cells must match the CSV, including literal quotes/newlines.
wb = load_workbook(out / "resolved.xlsx", data_only=True)
assert wb.sheetnames == ["Sheet1"]
values = list(wb.active.values)
assert list(values[0]) == list(table[0])
for row, cells in zip(table, values[1:]):
    assert (cells[1] or "") == row["cohort_label"]
    assert (cells[3] or "") == row["exclusion_label"]
    assert cells[2] == int(row["n_remaining"])
assert len(values) == len(table) + 1
# Actual renderer text/numeric contents and raster decoding, not existence.
boxes = [x["box"].text for x in d._layout() if x["type"] == "main"]
assert str(expected_n) in boxes[-1]
assert table[-1]["cohort_label"] in boxes[-1]
fig, ax = d.render()
assert boxes[-1] in [text.get_text() for text in ax.texts]
m.plt.close(fig)
im = Image.open(out / "figure.png").convert("RGB")
assert im.width > 100 and im.height > 100
assert len(im.getcolors(maxcolors=im.width * im.height)) > 10
print("ARTIFACT_CONTENT_PASS", mode)
